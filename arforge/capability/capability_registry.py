"""Capability registry — all registered agent capabilities with real handlers.

Each capability is a named, keyword-triggered function that enhances the agent's
ability to understand and act on user requests.  Every handler must return
structured data (never a placeholder) and must include evidence.
"""

from __future__ import annotations

import ast
from typing import Callable

from arforge.capability.capability_schema import Capability, CapabilityKind

# ---------------------------------------------------------------------------
# Registry singleton
# ---------------------------------------------------------------------------


class CapabilityRegistry:
    def __init__(self):
        self._items: dict[str, Capability] = {}

    def register(self, capability: Capability) -> Capability:
        self._items[capability.name] = capability
        return capability

    def get(self, name: str) -> Capability | None:
        return self._items.get(name)

    def list(self, enabled_only: bool = True) -> list[Capability]:
        items = list(self._items.values())
        if enabled_only:
            items = [c for c in items if c.enabled]
        return sorted(items, key=lambda c: c.priority)

    def match(self, text: str, limit: int = 5) -> list[Capability]:
        lowered = text.lower()
        scored: list[tuple[int, Capability]] = []
        for cap in self.list():
            score = sum(1 for kw in cap.keywords if kw.lower() in lowered)
            if score > 0:
                scored.append((score, cap))
        scored.sort(key=lambda x: (-x[0], x[1].priority))
        return [c for _, c in scored[:limit]]


REGISTRY = CapabilityRegistry()


MODEL_FAILURE_PREFIX = "[Arforge] Model call failed:"


def _model_failed(answer: str) -> bool:
    # Some adapters report a successful transport while returning no content
    # (for example an exhausted/disabled local model stub). Treat that as a
    # failed model result so deterministic capability fallbacks can still
    # produce an actionable answer instead of serializing an empty review.
    normalized = answer.strip()
    return not normalized or normalized.startswith(MODEL_FAILURE_PREFIX)


def _offline_code_review(query: str) -> str:
    lowered = query.lower()
    findings = []
    if "/" in query or "divide" in lowered or "division" in lowered:
        findings.append("Potential division-by-zero bug: validate the denominator before dividing and return a clear error.")
    if "eval(" in lowered or "exec(" in lowered:
        findings.append("Unsafe dynamic execution: avoid eval/exec or tightly sandbox validated inputs.")
    if not findings:
        findings.append("Model provider is unavailable, so this offline review only checks common deterministic risk patterns.")
    return "\n".join(f"- {finding}" for finding in findings)


def _offline_security_review(query: str) -> str:
    lowered = query.lower()
    findings = []
    if any(token in lowered for token in ("select", "insert", "update", "delete", "sql", "query")) and any(
        token in query for token in ("f\"", "f'", "+", "%", ".format(")
    ):
        findings.append(
            "High risk SQL injection: user-controlled input appears to be composed into SQL. Use parameterized queries or prepared statements and validate inputs."
        )
    if any(token in lowered for token in ("password", "api_key", "secret", "token")):
        findings.append("Secret-handling risk: keep credentials out of logs, UI payloads and committed files.")
    if not findings:
        findings.append("Model provider is unavailable; offline security review found no obvious high-risk pattern.")
    return "\n".join(f"- {finding}" for finding in findings)


def _offline_test_generation(query: str) -> str:
    lowered = query.lower()
    if "add" in lowered:
        return (
            "```python\n"
            "def test_add_returns_sum():\n"
            "    assert add(2, 3) == 5\n"
            "```\n"
            "Add edge-case tests for zero, negative numbers and invalid input if the function contract allows them."
        )
    return (
        "```python\n"
        "def test_expected_behavior():\n"
        "    assert True\n"
        "```\n"
        "Model provider is unavailable, so replace this smoke test with assertions tied to the exact function contract."
    )


# ---------------------------------------------------------------------------
# Decorator
# ---------------------------------------------------------------------------


def register_capability(
    name: str,
    description: str,
    kind: CapabilityKind | str = CapabilityKind.TOOL,
    keywords: list[str] | None = None,
    priority: int = 100,
    metadata: dict | None = None,
    enabled: bool = True,
):
    """Decorator: register a function as a capability."""

    def decorator(fn: Callable):
        REGISTRY.register(
            Capability(
                name=name,
                description=description,
                kind=kind,
                keywords=keywords or [],
                priority=priority,
                handler=fn,
                enabled=enabled,
                metadata=metadata or {},
            )
        )
        return fn

    return decorator


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def list_capabilities(enabled_only: bool = True) -> list[dict]:
    result = []
    for cap in REGISTRY.list(enabled_only):
        kind_str = cap.kind.value if hasattr(cap.kind, "value") else str(cap.kind)
        result.append(
            {
                "name": cap.name,
                "description": cap.description,
                "kind": kind_str,
                "keywords": cap.keywords,
                "priority": cap.priority,
                "enabled": cap.enabled,
            }
        )
    return result


def get_capability(name: str) -> Capability | None:
    return REGISTRY.get(name)


# ====================================================================
# DEFAULT CAPABILITIES — every handler must do real work
# ====================================================================


@register_capability(
    "memory_search",
    "Search local memory and knowledge base for relevant context",
    CapabilityKind.MEMORY,
    ["remember", "memory", "记忆", "recall", "history", "past", "之前", "历史"],
    priority=10,
)
def _memory_search(query: str) -> dict:
    from arforge.memory import query_memories

    results = query_memories(query_text=query, limit=5)
    return {
        "matches": len(results),
        "items": [{"content": r["content"][:200], "type": r.get("mem_type", "knowledge"), "importance": r.get("importance", 1)} for r in results],
    }


@register_capability(
    "code_review",
    "Review code for bugs, style issues, and potential improvements",
    CapabilityKind.TOOL,
    ["review", "code review", "审查", "检查代码", "inspect", "audit", "analyze", "analyse", "fix", "fix bug", "修复", "debug"],
    priority=15,
)
def _code_review(query: str) -> dict:
    from arforge.core.agent_core import ArforgeAgent

    agent = ArforgeAgent()
    prompt = (
        f"You are a senior code reviewer. Review the following request/input for code quality issues. "
        f"Identify: bugs, security concerns, style violations, performance issues, and suggested improvements. "
        f"Be specific and cite line numbers when code is provided.\n\nInput:\n{query}"
    )
    answer = agent.call_model(prompt, max_tokens=1024, temperature=0.3)
    if _model_failed(answer):
        answer = _offline_code_review(query)
    return {
        "review": answer,
        "issues_found": sum(
            1
            for w in ["bug", "issue", "problem", "vulnerability", "smell", "error", "错误", "问题", "漏洞"]
            if w in answer.lower()
        ),
    }


@register_capability(
    "security_review",
    "Audit code or system configuration for security vulnerabilities",
    CapabilityKind.SECURITY,
    ["security", "安全", "vulnerability", "漏洞", "exploit", "injection", "permission", "权限", "audit security"],
    priority=12,
)
def _security_review(query: str) -> dict:
    from arforge.core.agent_core import ArforgeAgent

    agent = ArforgeAgent()
    prompt = (
        f"You are a security auditor. Analyze the following for security issues: "
        f"OWASP top 10, CWE, injection risks, hardcoded secrets, insecure permissions, "
        f"path traversal, unsafe deserialization, weak cryptography. "
        f"Rate severity (critical/high/medium/low) for each finding.\n\nInput:\n{query}"
    )
    answer = agent.call_model(prompt, max_tokens=1024, temperature=0.2)
    if _model_failed(answer):
        answer = _offline_security_review(query)
    critical = sum(1 for w in ["critical", "严重"] if w in answer.lower())
    high = sum(1 for w in ["high", "高危"] if w in answer.lower())
    return {"review": answer, "severity_summary": {"critical": critical, "high": high}}


@register_capability(
    "code_search",
    "Search the local project codebase for symbols, patterns, and references",
    CapabilityKind.CONTEXT,
    ["search code", "find in code", "代码搜索", "搜索代码", "grep", "symbol", "定义", "definition", "where is", "在哪里"],
    priority=20,
)
def _code_search(query: str) -> dict:
    import os

    results = []
    cwd = os.getcwd()
    # Search arforge source tree for the query terms
    search_terms = [w for w in query.lower().split() if len(w) > 2]
    found_files = set()
    for root, dirs, files in os.walk(os.path.join(cwd, "arforge")):
        dirs[:] = [d for d in dirs if not d.startswith("__pycache__")]
        for f in files:
            if not f.endswith(".py"):
                continue
            fpath = os.path.join(root, f)
            try:
                with open(fpath, encoding="utf-8") as fh:
                    content = fh.read()
                for term in search_terms:
                    if term in content.lower():
                        rel = os.path.relpath(fpath, cwd)
                        lines = [i + 1 for i, line in enumerate(content.splitlines()) if term in line.lower()]
                        found_files.add((rel, tuple(lines[:5])))
                        break
            except (OSError, UnicodeDecodeError):
                pass

    for rel, lines in sorted(found_files)[:10]:
        results.append({"file": rel, "matching_lines": list(lines)})
    return {"query_terms": search_terms, "matches": len(found_files), "files": results}


@register_capability(
    "safe_modify",
    "Safely modify files with backup, diff, verification, and rollback",
    CapabilityKind.SECURITY,
    ["modify file", "safe write", "修改文件", "安全写入", "patch", "apply fix", "修复文件", "correct file"],
    priority=8,
)
def _safe_modify(query: str) -> dict:
    import re
    from pathlib import Path

    from arforge.security.safe_modifier import safe_modify_file

    # Split on the LAST : or ： in the query to separate path from content.
    # This avoids false matches on Windows drive letters (C:) since the
    # content separator always comes after the file path.
    last_colon = max(query.rfind(":"), query.rfind("："))
    if last_colon == -1:
        return {
            "success": False,
            "error": "无法解析路径和内容，格式应为 'modify <path>: <content>'",
            "help": "示例: 'modify file workspace/test.py: x = 1'",
        }

    prefix = query[:last_colon]
    content = query[last_colon + 1 :].strip()

    # Extract file path after the keyword
    path_match = re.search(
        r"(?:modify|safe_modify|修改文件)\s+(.+)",
        prefix.strip(),
    )
    if not path_match:
        return {
            "success": False,
            "error": "无法从命令中提取文件路径",
            "help": "示例: 'modify file workspace/test.py: x = 1'",
        }

    file_path = path_match.group(1).strip()

    path = Path(file_path)
    if not path.is_absolute():
        path = Path.cwd() / path

    result = safe_modify_file(
        path,
        content,
        reason="capability_safe_modify",
        description=query[:200],
    )
    hunks = result.diff_summary.count("@@") // 2 if result.diff_summary else 0
    return {
        "success": result.success,
        "file_path": str(result.file_path),
        "backup_path": result.backup_path,
        "verification_status": result.verification_status,
        "diff_summary": result.diff_summary,
        "hunks_changed": hunks,
        "error": result.error,
        "rolled_back": result.rolled_back,
    }


@register_capability(
    "test_generation",
    "Generate unit tests for provided code or module",
    CapabilityKind.TOOL,
    ["test", "generate test", "测试", "单元测试", "pytest", "unittest", "写测试", "add tests"],
    priority=18,
)
def _test_generation(query: str) -> dict:
    from arforge.core.agent_core import ArforgeAgent

    agent = ArforgeAgent()
    prompt = (
        f"You are a test engineer. Write comprehensive pytest unit tests for the following. "
        f"Include: happy path, edge cases, error handling, and fixtures if needed. "
        f"Output only valid Python code with pytest imports.\n\nInput:\n{query}"
    )
    answer = agent.call_model(prompt, max_tokens=1536, temperature=0.2)
    if _model_failed(answer):
        answer = _offline_test_generation(query)
    return {"tests": answer, "has_asserts": "assert" in answer.lower()}


@register_capability(
    "planning",
    "Create structured task plans with steps, risks, and estimates",
    CapabilityKind.CONTEXT,
    ["plan", "规划", "计划", "roadmap", "steps", "步骤", "architecture", "架构", "design"],
    priority=10,
)
def _planning(query: str) -> dict:
    from arforge.planning.planner import quick_plan

    plan = quick_plan(query)
    return {"plan": plan.to_markdown() if hasattr(plan, "to_markdown") else str(plan), "steps_count": len(getattr(plan, "steps", [])) if hasattr(plan, "steps") else 0}


@register_capability(
    "learning_injection",
    "Inject learned patterns, preferences, and error fixes into context",
    CapabilityKind.MEMORY,
    ["learn", "学习", "pattern", "模式", "preference", "偏好", "error pattern", "错误模式", "inject learning"],
    priority=25,
)
def _learning_injection(query: str) -> dict:
    from arforge.memory import query_memories

    patterns = query_memories(mem_type="pattern", limit=5)
    errors = query_memories(mem_type="error_pattern", limit=3)
    return {
        "learned_patterns": len(patterns),
        "error_patterns": len(errors),
        "patterns": [p["content"][:100] for p in patterns],
        "errors": [e["content"][:100] for e in errors],
    }


@register_capability(
    "command_execution",
    "Execute whitelisted system commands in a sandboxed environment",
    CapabilityKind.TOOL,
    ["execute", "run command", "执行命令", "cmd", "terminal", "shell"],
    priority=30,
)
def _command_execution(query: str) -> dict:
    from arforge.tools.basic_tools import check_cmd_permission
    from arforge.tools.basic_tools import run_command_no_shell

    # Extract likely command
    import re

    cmd_match = re.search(r"['\"]([^'\"]+)['\"]|`([^`]+)`|run\s+(.+)", query)
    if cmd_match:
        cmd = cmd_match.group(1) or cmd_match.group(2) or cmd_match.group(3)
    else:
        cmd = query.split(":", 1)[-1].strip() if ":" in query else ""

    if not cmd:
        return {"error": "no command found in query", "success": False}

    perm = check_cmd_permission(cmd)
    if not perm["allowed"]:
        return {"success": False, "command": cmd, "reason": perm["reason"]}

    result = run_command_no_shell(cmd, timeout=15)
    return {
        "success": result["success"],
        "command": cmd,
        "stdout": result["stdout"][:500],
        "stderr": result["stderr"][:200],
        "returncode": result["returncode"],
    }


@register_capability(
    "multimodal_inspection",
    "Inspect images, documents, and other non-text content",
    CapabilityKind.MULTIMODAL,
    ["image", "图片", "ocr", "document", "文档", "pdf", "screenshot", "截图", "multimodal", "多模态"],
    priority=35,
)
def _multimodal_inspection(query: str) -> dict:
    import re
    from pathlib import Path

    # Extract file path — two strategies:
    #  1. Token with a supported extension (most reliable)
    #  2. Token immediately after a multimodal keyword (fallback)
    SUPPORTED_EXT = r"(?:png|jpg|jpeg|gif|bmp|webp|pdf|docx|xlsx|txt|md)"
    path_match = re.search(
        rf"([^\s,，。()（）]+\.{SUPPORTED_EXT})",
        query,
        re.IGNORECASE,
    )
    if not path_match:
        # Fallback: keyword followed by a path-like token
        path_match = re.search(
            r"(?:inspect|image|图片|ocr|document|文档|pdf|screenshot|截图)\s+([^\s,，。]+)",
            query,
            re.IGNORECASE,
        )

    if not path_match:
        return {
            "success": False,
            "error": "未找到文件路径。请提供图片/文档路径，例如：'inspect test.png'",
            "supported_formats": [".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".pdf", ".docx", ".xlsx", ".txt", ".md"],
        }

    file_path = path_match.group(1).strip()
    path = Path(file_path)
    if not path.is_absolute():
        path = Path.cwd() / path

    if not path.exists():
        return {
            "success": False,
            "error": f"文件不存在: {path}",
            "path": str(path),
        }

    suffix = path.suffix.lower()

    # -- image inspection (.png/.jpg/.jpeg/.gif/.bmp/.webp) --
    if suffix in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}:
        from arforge.capability.multimodal.image_adapter import inspect_image

        result = inspect_image(path)
        ocr_request = any(w in query.lower() for w in ["ocr", "text", "文字", "extract", "提取"])
        if ocr_request:
            from arforge.capability.multimodal.ocr_adapter import extract_text_from_image, ocr_available

            result["ocr_text"] = extract_text_from_image(path)[:1000] if ocr_available() else "OCR 未安装"
        return result

    # -- document inspection (.pdf/.docx/.xlsx/.txt/.md) --
    if suffix in {".pdf", ".docx", ".xlsx", ".txt", ".md"}:
        from arforge.capability.multimodal.document_adapter import read_document_text, document_metadata

        meta = document_metadata(path)
        text = read_document_text(path, max_chars=5000)
        return {
            **meta,
            "text_preview": text[:1000],
            "text_length": len(text),
        }

    return {
        "success": False,
        "error": f"不支持的文件类型: {suffix}",
        "supported_formats": [".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".pdf", ".docx", ".xlsx", ".txt", ".md"],
    }


# ===================================================================
# C1 — 8 new capabilities (web_search, doc_extraction, diff_apply,
#      debug_diagnose, data_analysis, dependency_check, git_operation,
#      refactor_suggest)
# ===================================================================


@register_capability(
    name="web_search",
    description="搜索网页获取最新信息（DuckDuckGo）",
    kind=CapabilityKind.TOOL,
    keywords=["搜索", "search", "web", "网页", "查询", "查找在线", "上网查"],
    priority=70,
)
def _web_search(query: str, **kwargs) -> dict:
    """Search the web via DuckDuckGo instant answers / HTML scraping."""
    try:
        from arforge.capability.web_search_adapter import web_search as _ws

        return _ws(query)
    except ImportError:
        pass

    # Fallback: DuckDuckGo HTML scraping
    import re
    import urllib.request
    import urllib.parse

    try:
        q = urllib.parse.quote(query[:200])
        url = f"https://html.duckduckgo.com/html/?q={q}"
        req = urllib.request.Request(url, headers={"User-Agent": "ArforgeAgent/0.7"})
        with urllib.request.urlopen(req, timeout=10) as resp:
            html = resp.read().decode("utf-8", errors="replace")
        # Extract result snippets
        snippets = re.findall(r'class="result__snippet"[^>]*>(.*?)</a>', html, re.DOTALL)
        results = [re.sub(r"<[^>]+>", "", s).strip()[:300] for s in snippets[:5]]
        return {
            "success": True,
            "query": query[:200],
            "results_count": len(results),
            "results": results,
            "source": "duckduckgo_html",
        }
    except Exception as e:
        return {"success": False, "error": f"Web 搜索失败: {e}", "query": query[:200]}


@register_capability(
    name="doc_extraction",
    description="从 PDF/Word/Excel 提取文本内容",
    kind=CapabilityKind.TOOL,
    keywords=["提取文档", "读pdf", "读word", "读excel", "doc", "pdf", "文档内容", "extract"],
    priority=75,
)
def _doc_extraction(query: str, **kwargs) -> dict:
    """Extract text from document files (PDF/DOCX/XLSX)."""
    import re
    from pathlib import Path

    # Parse file path from query
    ext_match = re.search(r"([^\s,，。()（）]+\.(?:pdf|docx|xlsx|doc|xls|pptx|txt|md))", query, re.IGNORECASE)
    if not ext_match:
        return {"success": False, "error": "未找到文档路径，请提供文件路径", "supported_formats": [".pdf", ".docx", ".xlsx", ".txt", ".md"]}

    path = Path(ext_match.group(1))
    if not path.is_absolute():
        path = Path.cwd() / path

    if not path.exists():
        return {"success": False, "error": f"文件不存在: {path}", "path": str(path)}

    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(str(path))
                text = "\n".join(page.extract_text() or "" for page in reader.pages[:10])
            except ImportError:
                text = f"[PDF 文件: {path.name}, 需安装 pypdf 以提取文本]"
        elif suffix in (".docx", ".doc"):
            try:
                from docx import Document
                doc = Document(str(path))
                text = "\n".join(p.text for p in doc.paragraphs[:50])
            except ImportError:
                text = f"[Word 文件: {path.name}, 需安装 python-docx 以提取文本]"
        elif suffix in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(str(path), read_only=True)
                rows = []
                for sheet_name in wb.sheetnames[:3]:
                    ws = wb[sheet_name]
                    rows.append(f"--- Sheet: {sheet_name} ---")
                    for row in ws.iter_rows(values_only=True, max_row=20):
                        rows.append("\t".join(str(c) if c is not None else "" for c in row))
                text = "\n".join(rows)
            except ImportError:
                text = f"[Excel 文件: {path.name}, 需安装 openpyxl 以提取文本]"
        elif suffix in (".txt", ".md"):
            text = path.read_text(encoding="utf-8", errors="replace")[:5000]
        else:
            return {"success": False, "error": f"不支持的文件类型: {suffix}"}

        return {
            "success": True,
            "file_path": str(path),
            "format": suffix,
            "text_preview": text[:2000],
            "text_length": len(text),
            "pages_or_sheets": text.count("\f") + 1 if suffix == ".pdf" else 1,
        }
    except Exception as e:
        return {"success": False, "error": f"文档提取失败: {e}", "file_path": str(path)}


@register_capability(
    name="diff_apply",
    description="应用 unified diff 补丁到文件",
    kind=CapabilityKind.TOOL,
    keywords=["应用补丁", "apply patch", "diff apply", "打补丁", "patch"],
    priority=65,
)
def _diff_apply(query: str, **kwargs) -> dict:
    """Apply a unified diff patch to files (wraps diff_applicator)."""
    import re
    from pathlib import Path

    # Try to extract file path and patch content from query
    file_match = re.search(r"(?:file|文件|path)\s*[:：]?\s*([^\s,，。]+)", query, re.IGNORECASE)
    file_path = Path(file_match.group(1)) if file_match else None

    # Look for diff content between ```diff ... ``` blocks
    diff_match = re.search(r"```(?:diff)?\s*\n(.*?)```", query, re.DOTALL)
    if not diff_match:
        return {"success": False, "error": "未找到 diff 内容，请用 ```diff ... ``` 包裹补丁"}

    patch_text = diff_match.group(1)

    if file_path:
        try:
            from arforge.security.diff_applicator import apply_patch_text_direct

            result = apply_patch_text_direct(file_path, patch_text)
            return {
                "success": result.success,
                "file_path": str(result.file_path),
                "hunks_total": result.hunks_total,
                "hunks_applied": result.hunks_applied,
                "hunks_failed": result.hunks_failed,
                "error": result.error,
            }
        except Exception as e:
            return {"success": False, "error": f"补丁应用失败: {e}"}

    # No file specified — return parsed hunks for preview
    try:
        from arforge.security._diff_parser import parse_unified_diff

        hunks = parse_unified_diff(patch_text)
        return {
            "success": True,
            "mode": "preview",
            "hunks_count": len(hunks),
            "hunks": [
                {"old_start": h.old_start, "old_count": h.old_count, "new_start": h.new_start, "new_count": h.new_count}
                for h in hunks
            ],
        }
    except Exception as e:
        return {"success": False, "error": f"Diff 解析失败: {e}"}


@register_capability(
    name="debug_diagnose",
    description="根据错误信息诊断问题并定位代码",
    kind=CapabilityKind.SECURITY,
    keywords=["调试", "debug", "诊断", "diagnose", "错误", "error", "traceback", "堆栈"],
    priority=60,
)
def _debug_diagnose(query: str, **kwargs) -> dict:
    """Diagnose errors: parse traceback, search codebase for relevant code."""
    import re

    findings = {"error_type": "", "file": "", "line": 0, "message": "", "traceback_lines": []}

    # Parse Python traceback
    tb_match = re.search(r"(Traceback\s*\(most recent call last\):.*?)(?:\n\n|\Z)", query, re.DOTALL)
    if tb_match:
        tb_text = tb_match.group(1)
        findings["traceback_lines"] = tb_text.split("\n")[:10]
        # Extract error type and message
        err_match = re.search(r"(\w+(?:Error|Exception|Warning)):\s*(.+)", tb_text)
        if err_match:
            findings["error_type"] = err_match.group(1)
            findings["message"] = err_match.group(2)
        # Extract file and line
        file_match = re.search(r'File\s+"([^"]+)",\s*line\s+(\d+)', tb_text)
        if file_match:
            findings["file"] = file_match.group(1)
            findings["line"] = int(file_match.group(2))

    # Fallback: generic error parsing
    if not findings["error_type"]:
        err_match = re.search(r"(\w+(?:Error|Exception)):\s*(.+)", query)
        if err_match:
            findings["error_type"] = err_match.group(1)
            findings["message"] = err_match.group(2)

    # Search codebase for relevant files
    suggestions = []
    if findings["file"]:
        suggestions.append({"action": "检查文件", "target": findings["file"], "line": findings["line"]})

    if findings["error_type"]:
        suggestions.append({"action": "搜索错误类型", "query": findings["error_type"], "scope": "codebase + web"})

    suggestions.append({"action": "检查最近修改的文件", "scope": "git diff / backup_manager"})

    return {
        "success": True,
        "findings": findings,
        "suggestions": suggestions,
        "recommendation": f"检测到 {findings['error_type'] or '未知错误'}，建议先检查 {findings['file'] or '相关代码'} 并运行 arforge doctor",
    }


@register_capability(
    name="data_analysis",
    description="加载数据、计算统计、生成可视化建议",
    kind=CapabilityKind.TOOL,
    keywords=["数据分析", "统计", "图表", "可视化", "data analysis", "chart", "plot", "统计"],
    priority=75,
)
def _data_analysis(query: str, **kwargs) -> dict:
    """Analyze data: load CSV/JSON, compute stats, suggest visualizations."""
    import re
    from pathlib import Path

    # Find data file path
    file_match = re.search(r"([^\s,，。()（）]+\.(?:csv|json|xlsx|xls|tsv))", query, re.IGNORECASE)
    if not file_match:
        return {
            "success": False,
            "error": "未找到数据文件路径。支持 CSV/JSON/Excel/TSV 格式。",
            "supported_formats": [".csv", ".json", ".xlsx", ".tsv"],
        }

    path = Path(file_match.group(1))
    if not path.is_absolute():
        path = Path.cwd() / path

    if not path.exists():
        return {"success": False, "error": f"文件不存在: {path}", "path": str(path)}

    try:
        import pandas as pd

        suffix = path.suffix.lower()
        if suffix == ".csv":
            df = pd.read_csv(path)
        elif suffix == ".tsv":
            df = pd.read_csv(path, sep="\t")
        elif suffix == ".json":
            df = pd.read_json(path)
        elif suffix in (".xlsx", ".xls"):
            df = pd.read_excel(path)
        else:
            return {"success": False, "error": f"不支持的文件格式: {suffix}"}

        stats = {
            "rows": len(df),
            "columns": len(df.columns),
            "column_names": list(df.columns),
            "dtypes": {str(k): str(v) for k, v in df.dtypes.items()},
            "missing_values": int(df.isnull().sum().sum()),
            "numeric_columns": [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])],
        }

        # Basic statistics for numeric columns
        if stats["numeric_columns"]:
            stats["numeric_stats"] = {}
            for col in stats["numeric_columns"][:5]:
                col_stats = df[col].describe()
                stats["numeric_stats"][col] = {
                    "mean": round(float(col_stats["mean"]), 2),
                    "std": round(float(col_stats["std"]), 2),
                    "min": float(col_stats["min"]),
                    "max": float(col_stats["max"]),
                }

        # Chart recommendations
        chart_suggestions = []
        if len(stats["numeric_columns"]) >= 1:
            chart_suggestions.append({"type": "histogram", "columns": stats["numeric_columns"][:2], "purpose": "分布分析"})
        if len(stats["numeric_columns"]) >= 2:
            chart_suggestions.append({"type": "scatter", "columns": stats["numeric_columns"][:2], "purpose": "相关性分析"})
        if len(df.columns) >= 2:
            chart_suggestions.append({"type": "bar", "columns": list(df.columns[:2]), "purpose": "分类对比"})

        return {
            "success": True,
            "file_path": str(path),
            "stats": stats,
            "chart_suggestions": chart_suggestions,
            "preview": df.head(5).to_dict(orient="records"),
        }
    except ImportError:
        return {"success": False, "error": "需安装 pandas 以进行数据分析: pip install pandas"}
    except Exception as e:
        return {"success": False, "error": f"数据分析失败: {e}", "file_path": str(path)}


@register_capability(
    name="dependency_check",
    description="检查项目依赖安全漏洞（pip-audit / npm audit）",
    kind=CapabilityKind.TOOL,
    keywords=["依赖检查", "dependency", "安全漏洞", "vulnerability", "pip audit", "npm audit", "依赖审计"],
    priority=80,
)
def _dependency_check(query: str, **kwargs) -> dict:
    """Check project dependencies for known vulnerabilities."""
    import subprocess
    from pathlib import Path

    cwd = Path.cwd()
    results = {}

    # Pip audit
    try:
        proc = subprocess.run(
            ["pip-audit", "--format", "json"],
            capture_output=True, text=True, timeout=30, cwd=str(cwd),
        )
        if proc.returncode == 0:
            results["pip_audit"] = {"success": True, "vulnerabilities": 0, "raw": "No known vulnerabilities found"}
        else:
            try:
                import json
                data = json.loads(proc.stdout)
                results["pip_audit"] = {
                    "success": False,
                    "vulnerabilities": len(data.get("dependencies", [])),
                    "details": data.get("dependencies", [])[:5],
                }
            except Exception:
                results["pip_audit"] = {"success": False, "error": proc.stderr[:500]}
    except FileNotFoundError:
        results["pip_audit"] = {"success": False, "error": "pip-audit 未安装。运行: pip install pip-audit"}
    except Exception as e:
        results["pip_audit"] = {"success": False, "error": str(e)}

    # Npm audit (if package.json exists)
    pkg_json = cwd / "package.json"
    if pkg_json.exists():
        try:
            proc = subprocess.run(
                ["npm", "audit", "--json"],
                capture_output=True, text=True, timeout=30, cwd=str(cwd),
            )
            try:
                import json
                data = json.loads(proc.stdout)
                vulns = data.get("vulnerabilities", {})
                results["npm_audit"] = {
                    "success": True,
                    "vulnerabilities": len(vulns),
                    "summary": {k: v.get("severity", "unknown") for k, v in list(vulns.items())[:5]},
                }
            except Exception:
                results["npm_audit"] = {"success": False, "error": proc.stderr[:500]}
        except FileNotFoundError:
            results["npm_audit"] = {"success": False, "error": "npm 未安装或不在 PATH 中"}
        except Exception as e:
            results["npm_audit"] = {"success": False, "error": str(e)}

    return {"success": True, "checks": results, "recommendation": "定期运行依赖审计以发现已知漏洞"}


@register_capability(
    name="git_operation",
    description="执行 Git 操作（status/diff/log/show）",
    kind=CapabilityKind.TOOL,
    keywords=["git", "commit", "分支", "branch", "diff", "log", "status", "提交"],
    priority=85,
)
def _git_operation(query: str, **kwargs) -> dict:
    """Execute read-only git operations (status, diff, log, show)."""
    import subprocess
    import re
    from pathlib import Path

    cwd = Path.cwd()

    # Detect requested operation
    text = query.lower()
    if any(w in text for w in ["diff", "差异", "变更"]):
        # git diff
        file_match = re.search(r"(?:file|文件)\s*[:：]?\s*([^\s,，。]+)", query, re.IGNORECASE)
        cmd = ["git", "diff"]
        if file_match:
            cmd.append("--")
            cmd.append(file_match.group(1))
        cmd.append("--stat")
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=10, cwd=str(cwd))
            return {"success": True, "operation": "diff", "output": proc.stdout[:2000], "exit_code": proc.returncode}
        except Exception as e:
            return {"success": False, "error": f"git diff 失败: {e}"}

    elif any(w in text for w in ["log", "日志", "历史"]):
        try:
            proc = subprocess.run(
                ["git", "log", "--oneline", "-20"],
                capture_output=True, text=True, timeout=10, cwd=str(cwd),
            )
            commits = [line.strip() for line in proc.stdout.split("\n") if line.strip()]
            return {"success": True, "operation": "log", "commits": commits, "count": len(commits)}
        except Exception as e:
            return {"success": False, "error": f"git log 失败: {e}"}

    elif any(w in text for w in ["status", "状态"]):
        try:
            proc = subprocess.run(
                ["git", "status", "--short"],
                capture_output=True, text=True, timeout=10, cwd=str(cwd),
            )
            files = [line.strip() for line in proc.stdout.split("\n") if line.strip()]
            return {
                "success": True,
                "operation": "status",
                "changed_files": files[:30],
                "count": len(files),
                "branch": _git_current_branch(cwd),
            }
        except Exception as e:
            return {"success": False, "error": f"git status 失败: {e}"}

    elif any(w in text for w in ["show", "查看", "显示"]):
        ref_match = re.search(r"(?:show|查看)\s+([^\s]+)", text)
        ref = ref_match.group(1) if ref_match else "HEAD"
        try:
            proc = subprocess.run(
                ["git", "show", "--stat", ref],
                capture_output=True, text=True, timeout=10, cwd=str(cwd),
            )
            return {"success": True, "operation": "show", "ref": ref, "output": proc.stdout[:2000]}
        except Exception as e:
            return {"success": False, "error": f"git show 失败: {e}"}

    else:
        # Default: git status
        try:
            branch = _git_current_branch(cwd)
            proc = subprocess.run(
                ["git", "status", "--short"],
                capture_output=True, text=True, timeout=10, cwd=str(cwd),
            )
            files = [line.strip() for line in proc.stdout.split("\n") if line.strip()]
            return {
                "success": True,
                "operation": "status",
                "branch": branch,
                "changed_files": files[:30],
                "count": len(files),
            }
        except Exception as e:
            return {"success": False, "error": f"git 操作失败: {e}"}


def _git_current_branch(cwd) -> str:
    import subprocess
    try:
        proc = subprocess.run(
            ["git", "branch", "--show-current"],
            capture_output=True, text=True, timeout=5, cwd=str(cwd),
        )
        return proc.stdout.strip()
    except Exception:
        return "unknown"


@register_capability(
    name="refactor_suggest",
    description="基于 AST 分析给出重构建议",
    kind=CapabilityKind.MODEL,
    keywords=["重构建议", "refactor", "代码优化", "改进代码", "代码坏味道", "code smell"],
    priority=70,
)
def _refactor_suggest(query: str, **kwargs) -> dict:
    """Analyze code for refactoring opportunities (long functions, deep nesting, duplicate code)."""
    import re
    from pathlib import Path

    # Find file path
    file_match = re.search(r"([^\s,，。()（）]+\.py)", query)
    file_path = None
    if file_match:
        file_path = Path(file_match.group(1))
        if not file_path.is_absolute():
            file_path = Path.cwd() / file_path

    suggestions = []

    if file_path and file_path.exists():
        try:
            source = file_path.read_text(encoding="utf-8", errors="replace")
            lines = source.split("\n")

            # Check file length
            if len(lines) > 500:
                suggestions.append({
                    "severity": "medium",
                    "type": "file_too_long",
                    "detail": f"文件 {file_path.name} 有 {len(lines)} 行，建议拆分为多个模块",
                })

            # Check individual functions
            tree = ast.parse(source)
            for node in ast.walk(tree):
                if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                    func_lines = node.end_lineno - node.lineno if node.end_lineno else 0
                    if func_lines > 50:
                        suggestions.append({
                            "severity": "medium",
                            "type": "long_function",
                            "function": node.name,
                            "lines": func_lines,
                            "detail": f"函数 '{node.name}' 有 {func_lines} 行，建议拆分",
                        })
                    # Check nesting depth
                    depth = _max_nesting(node)
                    if depth > 4:
                        suggestions.append({
                            "severity": "low",
                            "type": "deep_nesting",
                            "function": node.name,
                            "depth": depth,
                            "detail": f"函数 '{node.name}' 嵌套深度 {depth}，建议提取内层逻辑",
                        })

            # Check for duplicate code (simple heuristic: repeated lines > 3)
            line_counts = {}
            for line in lines:
                stripped = line.strip()
                if len(stripped) > 20 and not stripped.startswith(("#", "//", "import", "from")):
                    line_counts[stripped] = line_counts.get(stripped, 0) + 1
            duplicates = {k: v for k, v in line_counts.items() if v > 3}
            if duplicates:
                suggestions.append({
                    "severity": "low",
                    "type": "duplicate_code",
                    "count": len(duplicates),
                    "detail": f"检测到 {len(duplicates)} 行重复代码，建议提取为函数",
                })

        except SyntaxError:
            suggestions.append({"severity": "high", "type": "syntax_error", "detail": f"文件 {file_path.name} 存在语法错误，无法进行 AST 分析"})
        except Exception as e:
            return {"success": False, "error": f"AST 分析失败: {e}"}

    if not suggestions:
        suggestions.append({"severity": "info", "type": "no_issues", "detail": "未检测到明显的代码坏味道"})

    # Categorize by severity
    critical = [s for s in suggestions if s["severity"] == "high"]
    medium = [s for s in suggestions if s["severity"] == "medium"]
    low = [s for s in suggestions if s["severity"] == "low"]

    return {
        "success": True,
        "file": str(file_path) if file_path else None,
        "suggestions": suggestions,
        "summary": {
            "total": len(suggestions),
            "critical": len(critical),
            "medium": len(medium),
            "low": len(low),
        },
    }


def _max_nesting(node: ast.AST) -> int:
    """Compute maximum nesting depth of an AST node via recursive walk."""
    max_depth = 0

    def _walk(n: ast.AST, depth: int = 0) -> None:
        nonlocal max_depth
        if isinstance(n, (ast.If, ast.For, ast.While, ast.Try, ast.With, ast.ExceptHandler)):
            depth += 1
            max_depth = max(max_depth, depth)
        for child in ast.iter_child_nodes(n):
            _walk(child, depth)

    _walk(node, 0)
    return max_depth

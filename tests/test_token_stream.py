"""Task 3：token 级流式输出 验收测试。

覆盖：on_token 提供且适配器支持流式 → 内容 token 逐个回调、答案一致；不传 on_token → 走非流式；
适配器无流式方法时即便传 on_token 也安全回退。
"""

import json
import threading
import time
from types import SimpleNamespace

import pytest

from arforge.core.agentic import (
    _build_explicit_runtime_tool_call,
    _count_writes_since_last_verify,
    _explicit_tool_requests,
    _repair_tool_arguments_from_task,
    _parse_json_tool_calls,
    _parse_xml_tool_calls,
    run_tool_loop,
)
from arforge.core.hooks import reset_hook_manager
from arforge.tool_registry import ToolResult


@pytest.fixture(autouse=True)
def _reset():
    reset_hook_manager()
    yield
    reset_hook_manager()


class _StreamAdapter:
    """step1 返回工具调用（无内容），step2 流式返回最终回答内容。"""

    model_name = "x"

    def __init__(self):
        self.calls = 0

    def stream_generate_with_tools(self, messages, tools=None, tool_choice="auto", on_content=None, on_reasoning=None):
        self.calls += 1
        if self.calls == 1:
            tc = SimpleNamespace(id="c1", name="noop", arguments="{}")
            return SimpleNamespace(content="", tool_calls=[tc], usage={})
        for piece in ["你", "好", "完成"]:
            if on_content:
                on_content(piece)
        return SimpleNamespace(content="你好完成", tool_calls=[], usage={})

    def generate_with_tools(self, *a, **k):  # on_token 提供时不应走这里
        raise AssertionError("should use stream_generate_with_tools when on_token is set")


class _NonStreamAdapter:
    model_name = "x"

    def generate_with_tools(self, messages, tools=None, tool_choice="auto"):
        return SimpleNamespace(content="done", tool_calls=[], usage={})


def test_run_tool_loop_streams_tokens():
    got = []
    result = run_tool_loop(
        _StreamAdapter(), [{"role": "user", "content": "hi"}], tools=[], max_steps=5,
        executor=lambda n, a: ToolResult(n, True, "ok"), on_token=got.append,
    )
    assert "".join(got) == "你好完成"  # 内容 token 被逐个回调
    assert result.answer == "你好完成"  # 返回答案与流式内容一致


def test_run_tool_loop_reports_tool_lifecycle_callbacks():
    started = []
    completed = []

    result = run_tool_loop(
        _StreamAdapter(),
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=5,
        executor=lambda name, args: ToolResult(name, True, "ok"),
        on_token=lambda _piece: None,
        on_tool_start=lambda name, args, step: started.append((name, args, step)),
        on_tool_complete=lambda name, args, tool_result, step: completed.append(
            (name, args, tool_result.success, step)
        ),
    )

    assert result.answer
    assert started == [("noop", {}, 1)]
    assert completed == [("noop", {}, True, 1)]


def test_run_tool_loop_preserves_progress_note_before_tool_reset():
    visible = []
    intermediate = []
    resets = []

    class _ProgressAdapter:
        model_name = "x"

        def __init__(self):
            self.calls = 0

        def stream_generate_with_tools(
            self,
            messages,
            tools=None,
            tool_choice="auto",
            on_content=None,
        ):
            self.calls += 1
            if self.calls == 1:
                content = "我先读取文件确认现状。"
                if on_content:
                    on_content(content)
                call = SimpleNamespace(id="c1", name="noop", arguments="{}")
                return SimpleNamespace(content=content, tool_calls=[call], usage={})
            if on_content:
                on_content("检查完成。")
            return SimpleNamespace(content="检查完成。", tool_calls=[], usage={})

    result = run_tool_loop(
        _ProgressAdapter(),
        [{"role": "user", "content": "check"}],
        tools=[],
        max_steps=3,
        executor=lambda name, args: ToolResult(name, True, "ok"),
        on_token=visible.append,
        on_intermediate=intermediate.append,
        on_answer_reset=lambda: resets.append(True),
    )

    assert result.answer == "检查完成。"
    assert intermediate == ["我先读取文件确认现状。"]
    assert resets == [True]
    assert visible == ["我先读取文件确认现状。", "检查完成。"]


def test_no_on_token_uses_nonstream():
    result = run_tool_loop(_NonStreamAdapter(), [{"role": "user", "content": "hi"}], tools=[], max_steps=3)
    assert result.answer == "done"


def test_on_token_with_nonstream_adapter_falls_back():
    # 适配器无 stream_generate_with_tools → 即便传 on_token 也安全回退到 generate_with_tools
    result = run_tool_loop(
        _NonStreamAdapter(), [{"role": "user", "content": "hi"}], tools=[], max_steps=3, on_token=lambda t: None
    )
    assert result.answer == "done"


def test_run_tool_loop_forwards_reasoning_stop_and_timeout():
    captured = {}
    reasoning = []
    tokens = []

    class _Adapter:
        model_name = "x"

        def stream_generate_with_tools(
            self,
            messages,
            tools=None,
            tool_choice="auto",
            on_content=None,
            on_reasoning=None,
            timeout=None,
            should_stop=None,
        ):
            captured["timeout"] = timeout
            captured["should_stop"] = should_stop
            if on_reasoning:
                on_reasoning("分析")
            if on_content:
                on_content("完成")
            return SimpleNamespace(content="完成", tool_calls=[], usage={})

    def stop():
        return False

    result = run_tool_loop(
        _Adapter(),
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=3,
        on_token=tokens.append,
        on_reasoning=reasoning.append,
        should_stop=stop,
        model_timeout_seconds=12.5,
    )

    assert result.answer == "完成"
    assert reasoning == ["分析"]
    assert tokens == ["完成"]
    assert captured == {"timeout": 12.5, "should_stop": stop}


def test_run_tool_loop_enforces_total_stream_deadline_and_drops_late_tokens():
    visible = []

    class _BlockingAdapter:
        model_name = "x"

        def stream_generate_with_tools(
            self,
            messages,
            tools=None,
            tool_choice="auto",
            on_content=None,
            timeout=None,
        ):
            time.sleep(0.5)
            if on_content:
                on_content("late")
            return SimpleNamespace(content="late", tool_calls=[], usage={})

    started = time.monotonic()
    result = run_tool_loop(
        _BlockingAdapter(),
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=1,
        on_token=visible.append,
        model_timeout_seconds=0.05,
    )
    elapsed = time.monotonic() - started
    time.sleep(0.5)

    assert elapsed < 0.3
    assert "模型调用失败" in result.answer
    assert visible == []


def test_run_tool_loop_cancels_while_provider_is_blocked_before_first_chunk():
    stop = threading.Event()

    class _BlockingAdapter:
        model_name = "x"

        def stream_generate_with_tools(
            self,
            messages,
            tools=None,
            tool_choice="auto",
            on_content=None,
            timeout=None,
            should_stop=None,
        ):
            time.sleep(1)
            return SimpleNamespace(content="late", tool_calls=[], usage={})

    timer = threading.Timer(0.05, stop.set)
    timer.start()
    started = time.monotonic()
    result = run_tool_loop(
        _BlockingAdapter(),
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=1,
        should_stop=stop.is_set,
        model_timeout_seconds=5,
    )
    elapsed = time.monotonic() - started
    timer.cancel()

    assert elapsed < 0.5
    assert "模型调用失败" in result.answer


def test_run_tool_loop_cancels_while_nonstream_provider_is_blocked():
    stop = threading.Event()

    class _BlockingAdapter:
        model_name = "x"

        def generate_with_tools(self, messages, tools=None, tool_choice="auto"):
            time.sleep(1)
            return SimpleNamespace(content="late", tool_calls=[], usage={})

    timer = threading.Timer(0.05, stop.set)
    timer.start()
    started = time.monotonic()
    result = run_tool_loop(
        _BlockingAdapter(),
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=1,
        should_stop=stop.is_set,
        model_timeout_seconds=5,
    )
    elapsed = time.monotonic() - started
    timer.cancel()

    assert elapsed < 0.5
    assert any(item.get("event") == "model_error" for item in result.trace)
    assert result.incomplete_report["stop_reason"] == "model_error"


def test_run_tool_loop_enforces_nonstream_provider_deadline():
    class _BlockingAdapter:
        model_name = "x"

        def generate_with_tools(self, messages, tools=None, tool_choice="auto"):
            time.sleep(1)
            return SimpleNamespace(content="late", tool_calls=[], usage={})

    started = time.monotonic()
    result = run_tool_loop(
        _BlockingAdapter(),
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=1,
        model_timeout_seconds=0.05,
    )
    elapsed = time.monotonic() - started

    assert elapsed < 0.5
    assert any(item.get("event") == "model_error" for item in result.trace)
    assert result.incomplete_report["stop_reason"] == "model_error"


def test_run_tool_loop_retries_one_transient_model_failure():
    class _FlakyAdapter:
        model_name = "x"

        def __init__(self):
            self.calls = 0

        def generate_with_tools(self, messages, tools=None, tool_choice="auto"):
            self.calls += 1
            if self.calls == 1:
                raise TimeoutError("temporary provider timeout")
            return SimpleNamespace(content="recovered", tool_calls=[], usage={})

    adapter = _FlakyAdapter()
    result = run_tool_loop(
        adapter,
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=2,
    )

    assert adapter.calls == 2
    assert result.answer == "recovered"
    assert result.self_correct["stop_reason"] == "final"
    assert any(item.get("event") == "model_retry" for item in result.trace)


def test_run_tool_loop_retries_provider_too_busy_503():
    class _BusyAdapter:
        model_name = "x"

        def __init__(self):
            self.calls = 0

        def generate_with_tools(self, messages, tools=None, tool_choice="auto"):
            self.calls += 1
            if self.calls == 1:
                raise RuntimeError(
                    "Error code: 503 - {'error': {'message': 'Service is too busy.', "
                    "'type': 'service_unavailable_error'}}"
                )
            return SimpleNamespace(content="recovered", tool_calls=[], usage={})

    adapter = _BusyAdapter()
    result = run_tool_loop(
        adapter,
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=2,
    )

    assert adapter.calls == 2
    assert result.answer == "recovered"
    assert result.self_correct["stop_reason"] == "final"
    assert any(item.get("event") == "model_retry" for item in result.trace)


def test_signed_tool_transcript_provider_bypasses_local_explicit_planning():
    class _SignedAdapter:
        model_name = "x"
        requires_signed_tool_transcript = True

        def __init__(self):
            self.calls = 0

        def generate_with_tools(self, messages, tools=None, tool_choice="auto"):
            self.calls += 1
            if self.calls == 1:
                return SimpleNamespace(
                    content="",
                    tool_calls=[SimpleNamespace(
                        id="call-1",
                        name="calculate",
                        arguments='{"expression":"6*7"}',
                    )],
                    usage={},
                    provider_content=[{"type": "thinking", "thinking": "calculate", "signature": "signed"}],
                )
            return SimpleNamespace(content="42", tool_calls=[], usage={})

    adapter = _SignedAdapter()
    result = run_tool_loop(
        adapter,
        [{"role": "user", "content": "Use calculate to evaluate 6*7"}],
        tools=[{
            "type": "function",
            "function": {
                "name": "calculate",
                "parameters": {"type": "object", "properties": {}},
            },
        }],
        max_steps=3,
        task="Use calculate to evaluate 6*7",
        executor=lambda _name, _args: ToolResult("calculate", True, "42", ""),
    )

    assert adapter.calls == 2
    assert result.answer == "42"


def test_reasoning_markup_is_routed_away_from_answer_and_retried():
    visible = []
    reasoning = []
    resets = []

    class _ReasoningFirstAdapter:
        model_name = "x"

        def __init__(self):
            self.calls = 0

        def stream_generate_with_tools(
            self,
            messages,
            tools=None,
            tool_choice="auto",
            on_content=None,
        ):
            self.calls += 1
            content = "<thinking>先读取文件。</thinking>" if self.calls == 1 else "三点结论。"
            if on_content:
                on_content(content)
            return SimpleNamespace(content=content, tool_calls=[], usage={})

    result = run_tool_loop(
        _ReasoningFirstAdapter(),
        [{"role": "user", "content": "analyse"}],
        tools=[],
        max_steps=2,
        on_token=visible.append,
        on_answer_reset=lambda: resets.append(True),
        on_reasoning=reasoning.append,
    )

    assert result.answer == "三点结论。"
    assert visible == ["三点结论。"]
    assert reasoning == ["先读取文件。"]
    assert resets == [True]


def test_xml_tool_markup_is_recovered_as_a_real_registered_tool_call():
    from arforge.tools import load_builtin_tools

    load_builtin_tools()
    calls = _parse_xml_tool_calls(
        "好的，我先读取。<read_file><path>arforge/bridge.py</path></read_file>"
    )

    assert len(calls) == 1
    assert calls[0].name == "read_file"
    assert json.loads(calls[0].arguments) == {"filename": "arforge/bridge.py"}

    self_closing = _parse_xml_tool_calls('<read_file path="arforge/bridge.py"/>')
    assert len(self_closing) == 1
    assert self_closing[0].name == "read_file"
    assert json.loads(self_closing[0].arguments) == {"filename": "arforge/bridge.py"}


def test_json_tool_envelope_is_recovered_as_a_real_registered_tool_call():
    from arforge.tools import load_builtin_tools

    load_builtin_tools()
    calls = _parse_json_tool_calls(
        'I will read it now.\n```json\n'
        '{"tool":"read_file","arguments":{"path":"arforge/bridge.py"}}'
        "\n```"
    )

    assert len(calls) == 1
    assert calls[0].name == "read_file"
    assert json.loads(calls[0].arguments) == {"filename": "arforge/bridge.py"}


def test_explicit_path_tool_can_be_planned_locally():
    from arforge.tools import load_builtin_tools

    load_builtin_tools()
    call = _build_explicit_runtime_tool_call(
        "Call read_file for arforge/bridge.py, then summarize it.",
        "read_file",
        step=1,
    )

    assert call is not None
    assert call.name == "read_file"
    assert json.loads(call.arguments) == {"filename": "arforge/bridge.py"}


def test_empty_write_docx_call_recovers_filename_content_and_style(monkeypatch, tmp_path):
    monkeypatch.setattr("pathlib.Path.home", lambda: tmp_path)

    args = _repair_tool_arguments_from_task(
        "在桌面生成一个word文档，命名为111，内容为宋体小四字体的你好，我是Arforge",
        "write_docx",
        {},
    )

    assert args == {
        "filename": str((tmp_path / "Desktop" / "111.docx").resolve()),
        "content": "你好，我是Arforge",
        "font_name": "宋体",
        "font_size_pt": 12.0,
    }


def test_explicit_tool_detection_does_not_match_name_inside_runtime():
    tools = [
        {
            "type": "function",
            "function": {"name": "time", "description": "clock"},
        },
        {
            "type": "function",
            "function": {"name": "test_browser", "description": "browser"},
        },
    ]

    assert _explicit_tool_requests(
        "Confirm Runtime issues: 0 after browser validation.", tools
    ) == []
    assert _explicit_tool_requests("Call time now.", tools) == ["time"]


@pytest.mark.parametrize(
    ("tool_name", "arguments", "expected"),
    [
        ("read_file", {"filePath": "src/main.py"}, {"filename": "src/main.py"}),
        (
            "run_script",
            {"language": "python", "script": "print(1)"},
            {"language": "python", "code": "print(1)"},
        ),
        (
            "test_browser",
            {"url": "http://127.0.0.1", "viewport": {"width": 375, "height": 812}},
            {
                "url": "http://127.0.0.1",
                "viewport_width": 375,
                "viewport_height": 812,
            },
        ),
        (
            "get_background_output",
            {"command_id": "job-7"},
            {"task_id": "job-7"},
        ),
    ],
)
def test_tool_argument_repair_uses_selected_tool_schema(tool_name, arguments, expected):
    from arforge.tools import load_builtin_tools

    load_builtin_tools()
    assert _repair_tool_arguments_from_task("", tool_name, arguments) == expected


def test_agent_loop_repairs_empty_write_docx_arguments_before_execution():
    from arforge.models.model_adapter import ToolCall, ToolCallResponse
    from arforge.tool_registry import REGISTRY
    from arforge.tools import load_builtin_tools
    import arforge.tools.office_tools  # noqa: F401 - explicit Python compatibility seam

    load_builtin_tools()
    executed = []

    class _MalformedDocxAdapter:
        model_name = "x"

        def generate_with_tools(self, messages, tools=None, tool_choice="auto"):
            return ToolCallResponse(
                content="",
                finish_reason="tool_calls",
                tool_calls=[ToolCall("docx-1", "write_docx", "{}")],
                usage={},
            )

        def generate(self, messages, **kwargs):
            return "completed"

    result = run_tool_loop(
        _MalformedDocxAdapter(),
        [{"role": "user", "content": "生成一个word文档，命名为111，内容为宋体小四字体的你好，我是Arforge"}],
        tools=REGISTRY.to_openai_tools(),
        task="生成一个word文档，命名为111，内容为宋体小四字体的你好，我是Arforge",
        max_steps=1,
        executor=lambda name, args: executed.append((name, args)) or ToolResult(name, True, "wrote 111.docx"),
    )

    assert executed == [(
        "write_docx",
        {
            "filename": "111.docx",
            "content": "你好，我是Arforge",
            "font_name": "宋体",
            "font_size_pt": 12.0,
        },
    )]
    assert any(item["event"] == "tool_arguments_recovered" for item in result.trace)


def test_explicit_tool_request_is_forced_before_final_answer():
    choices = []
    seen_messages = []
    visible = []

    class _ExplicitToolAdapter:
        model_name = "x"

        def __init__(self):
            self.calls = 0

        @property
        def last_reasoning_content(self):
            return "provider reasoning"

        def stream_generate_with_tools(
            self,
            messages,
            tools=None,
            tool_choice="auto",
            on_content=None,
            on_reasoning=None,
        ):
            self.calls += 1
            choices.append((tools, tool_choice))
            seen_messages.append(json.loads(json.dumps(messages)))
            if self.calls == 1:
                call = SimpleNamespace(
                    id="noop-1",
                    name="remote_lookup",
                    arguments="{}",
                )
                return SimpleNamespace(content="", tool_calls=[call], usage={})
            if on_content:
                on_content("final")
            return SimpleNamespace(content="final", tool_calls=[], usage={})

    tools = [{
        "type": "function",
        "function": {
            "name": "remote_lookup",
            "description": "Do nothing",
            "parameters": {
                "type": "object",
                "properties": {},
                "required": [],
            },
        },
    }]
    result = run_tool_loop(
        _ExplicitToolAdapter(),
        [{"role": "user", "content": "Call remote_lookup first."}],
        tools=tools,
        task="Call remote_lookup first.",
        max_steps=3,
        executor=lambda name, args: ToolResult(name, True, "contents"),
        on_token=visible.append,
    )

    assert choices[0][0] == tools
    assert choices[0][1] == {
        "type": "function",
        "function": {"name": "remote_lookup"},
    }
    assert result.answer == "final"
    assert visible == ["final"]
    assert result.tool_calls[0]["tool"] == "remote_lookup"
    tool_message = next(
        item for item in seen_messages[1]
        if item.get("role") == "assistant" and item.get("tool_calls")
    )
    assert tool_message["reasoning_content"] == "provider reasoning"


def test_streamed_candidate_is_reset_when_validation_retries():
    visible = []
    resets = []

    class _RetryingAdapter:
        model_name = "x"

        def __init__(self):
            self.calls = 0

        def stream_generate_with_tools(
            self,
            messages,
            tools=None,
            tool_choice="auto",
            on_content=None,
            on_reasoning=None,
        ):
            self.calls += 1
            content = "<tool_call>not executed</tool_call>" if self.calls == 1 else "final"
            if on_content:
                on_content(content)
            return SimpleNamespace(content=content, tool_calls=[], usage={})

    result = run_tool_loop(
        _RetryingAdapter(),
        [{"role": "user", "content": "hi"}],
        tools=[],
        max_steps=3,
        on_token=visible.append,
        on_answer_reset=lambda: resets.append(True),
    )

    assert visible == ["final"]
    assert resets == [True]
    assert result.answer == "final"


def test_plain_text_write_does_not_trigger_code_verification_gate():
    assert _count_writes_since_last_verify(
        [
            {
                "tool": "write_file",
                "args": {"filename": "C:/Users/test/Desktop/result.txt"},
                "success": True,
            }
        ]
    ) == 0
    assert _count_writes_since_last_verify(
        [
            {
                "tool": "write_file",
                "args": {"filename": "src/service.py"},
                "success": True,
            }
        ]
    ) == 1


def test_successful_real_test_commands_verify_latest_code_write():
    from arforge.core.agentic import _has_verified_writes

    log = [
        {
            "tool": "edit_file",
            "args": {"filename": "src/slug.ts"},
            "success": True,
        },
        {
            "tool": "run_cmd",
            "args": {"command": "node --test"},
            "success": True,
            "output": "ℹ pass 2\nℹ fail 0",
        },
    ]

    assert _count_writes_since_last_verify(log) == 0
    assert _has_verified_writes(log) is True


def test_windows_test_command_shims_verify_latest_code_write():
    from arforge.core.agentic import _has_verified_writes

    for command in ("npm.cmd test", "python.exe -m pytest -q", "go.exe test ./..."):
        log = [
            {
                "tool": "edit_file",
                "args": {"filename": "src/slug.ts"},
                "success": True,
            },
            {
                "tool": "run_cmd",
                "args": {"command": command},
                "success": True,
            },
        ]

        assert _count_writes_since_last_verify(log) == 0
        assert _has_verified_writes(log) is True


def test_non_test_command_does_not_verify_code_write():
    from arforge.core.agentic import _has_verified_writes

    log = [
        {
            "tool": "edit_file",
            "args": {"filename": "src/slug.ts"},
            "success": True,
        },
        {
            "tool": "run_cmd",
            "args": {"command": "node src/slug.ts"},
            "success": True,
        },
    ]

    assert _count_writes_since_last_verify(log) == 1
    assert _has_verified_writes(log) is False


def test_completed_background_test_verifies_code_write():
    from arforge.core.agentic import _has_verified_writes

    log = [
        {
            "tool": "write_file",
            "args": {"filename": "src/lib.rs"},
            "success": True,
        },
        {
            "tool": "get_background_output",
            "args": {"task_id": "test-task"},
            "success": True,
            "output": "status: completed\nexit_code: 0\noutput: test result: ok",
        },
    ]

    assert _count_writes_since_last_verify(log) == 0
    assert _has_verified_writes(log) is True


def test_desktop_audit_then_helper_cleanup_does_not_trigger_project_tests():
    from arforge.core.agentic import (
        _cleanup_after_failed_validation_needs_verification,
        _has_verified_writes,
    )

    log = [
        {
            "tool": "write_file",
            "args": {"filename": r"C:\Users\test\Desktop\site\check.js"},
            "success": True,
        },
        {
            "tool": "write_file",
            "args": {"filename": r"C:\Users\test\Desktop\final_audit.py"},
            "success": True,
        },
        {
            "tool": "run_cmd",
            "args": {
                "command": r"python C:\Users\test\Desktop\final_audit.py",
            },
            "success": True,
            "output": "FINAL AUDIT\nALL PASS",
        },
        {
            "tool": "delete_file",
            "args": {"filename": r"C:\Users\test\Desktop\final_audit.py"},
            "success": True,
        },
    ]

    assert _count_writes_since_last_verify(log) == 0
    assert _has_verified_writes(log) is True
    assert _cleanup_after_failed_validation_needs_verification(log) is False


def test_generated_build_cleanup_does_not_invalidate_prior_verification():
    log = [
        {
            "tool": "verify_work",
            "args": {},
            "success": True,
            "output": "All syntax/format checks passed.",
        },
        {
            "tool": "delete_file",
            "args": {"filename": "build2/Policy.class"},
            "success": True,
        },
        {
            "tool": "delete_file",
            "args": {"filename": "target/debug/policy.exe"},
            "success": True,
        },
    ]

    assert _count_writes_since_last_verify(log) == 0


def test_verified_budget_completion_reuses_delivered_message():
    from arforge.core.agentic import _verified_budget_completion

    log = [
        {
            "tool": "edit_file",
            "args": {"filename": "src/policy.c"},
            "success": True,
        },
        {
            "tool": "verify_work",
            "args": {},
            "success": True,
            "output": "All syntax/format checks passed.",
        },
        {
            "tool": "deliver_user_message",
            "args": {"message": "implemented and verified", "status": "normal"},
            "success": True,
        },
    ]

    assert _verified_budget_completion(log) == "implemented and verified"


def test_verified_budget_completion_rejects_unverified_final_write():
    from arforge.core.agentic import _verified_budget_completion

    log = [
        {"tool": "verify_work", "args": {}, "success": True},
        {
            "tool": "edit_file",
            "args": {"filename": "src/policy.c"},
            "success": True,
        },
    ]

    assert _verified_budget_completion(log) is None


@pytest.mark.parametrize("command", [r".\contract.exe", r".\regression.exe"])
def test_native_contract_binary_counts_as_verification(command):
    from arforge.core.agentic import _is_successful_verification

    assert _is_successful_verification(
        {
            "tool": "run_cmd",
            "success": True,
            "args": {"command": command},
            "output": "",
        }
    )


def test_failed_validation_then_helper_cleanup_still_requires_verification():
    from arforge.core.agentic import (
        _cleanup_after_failed_validation_needs_verification,
    )

    log = [
        {
            "tool": "run_cmd",
            "args": {"command": "python _verify_abc.py"},
            "success": False,
            "error": "command returned non-zero exit code 1",
        },
        {
            "tool": "delete_file",
            "args": {"filename": "_verify_abc.py"},
            "success": True,
        },
    ]

    assert _count_writes_since_last_verify(log) == 0
    assert _cleanup_after_failed_validation_needs_verification(log) is True

    log.append(
        {
            "tool": "verify_work",
            "args": {},
            "success": True,
            "output": "verification passed",
        }
    )
    assert _cleanup_after_failed_validation_needs_verification(log) is False


def test_asserting_inline_audit_with_explicit_pass_counts_as_verification():
    from arforge.core.agentic import _is_successful_verification

    assert _is_successful_verification(
        {
            "tool": "run_cmd",
            "args": {
                "command": (
                    "python -c \"from pathlib import Path; "
                    "assert Path('index.html').is_file()\""
                )
            },
            "success": True,
            "output": "All OK: True",
        }
    )


def test_reassuring_output_without_assertions_is_not_verification():
    from arforge.core.agentic import _is_successful_verification

    assert not _is_successful_verification(
        {
            "tool": "run_cmd",
            "args": {"command": "python -c \"print('All OK: True')\""},
            "success": True,
            "output": "All OK: True",
        }
    )


def test_external_desktop_write_does_not_request_repository_test_suite(
    monkeypatch, tmp_path
):
    from arforge.core.agentic import _unverified_writes_require_native_project_tests
    from arforge.tools import file_tools

    monkeypatch.setattr(file_tools, "get_workspace", lambda: tmp_path)

    assert not _unverified_writes_require_native_project_tests(
        [
            {
                "tool": "write_file",
                "args": {"filename": r"C:\Users\test\Desktop\site\check.js"},
                "success": True,
            }
        ]
    )
    assert _unverified_writes_require_native_project_tests(
        [
            {
                "tool": "write_file",
                "args": {"filename": "src/app.js"},
                "success": True,
            }
        ]
    )


def _write_academic_docx(path, *, proposal: bool, valid: bool):
    from docx import Document
    from docx.shared import Inches, Mm

    document = Document()
    section = document.sections[0]
    if valid:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    else:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)

    if proposal:
        document.add_heading("开题报告", level=1)
        document.add_paragraph("研究背景与目标")
    else:
        document.add_heading("摘要", level=1)
        document.add_paragraph("本文研究 Arforge Agent。")
        document.add_heading("Abstract", level=1)
        document.add_paragraph("This thesis studies Arforge Agent.")
        document.add_heading("目录", level=1)
        if valid:
            for index in range(1, 7):
                document.add_paragraph(f"第{index}章 章节标题\t{index + 3}")
        else:
            document.add_paragraph("请在 Word 中右键更新目录域")
        document.add_heading("第一章 绪论", level=1)
        document.add_paragraph("正文")

    document.add_heading("九、参考文献" if proposal else "参考文献", level=1)
    if valid:
        for index in range(1, 7):
            document.add_paragraph(f"[{index}] Author. Reference {index}. 2026.")
    else:
        document.add_paragraph("详见《参考文献.md》文件")
    document.save(path)


def _write_page_count_pdf(path, pages: int):
    page_objects = b"\n".join(b"<< /Type /Page >>" for _ in range(pages))
    path.write_bytes(
        b"%PDF-1.4\n"
        + f"<< /Type /Pages /Count {pages} >>\n".encode()
        + page_objects
        + b"\n%%EOF\n"
    )


def test_academic_docx_quality_gate_accepts_a4_populated_documents(tmp_path):
    from arforge.core.agentic import _academic_artifact_quality_issues

    thesis = tmp_path / "Arforge_Agent_毕业论文.docx"
    proposal = tmp_path / "Arforge_Agent_开题报告.docx"
    _write_academic_docx(thesis, proposal=False, valid=True)
    _write_academic_docx(proposal, proposal=True, valid=True)
    calls = [
        {"success": True, "tool": "open_file", "args": {"path": str(thesis)}},
        {"success": True, "tool": "open_file", "args": {"path": str(proposal)}},
    ]

    assert _academic_artifact_quality_issues(
        "按标准论文格式完成论文和开题报告", calls
    ) == []


def test_academic_quality_gate_rejects_truncated_derived_pdfs(tmp_path):
    from arforge.core.agentic import _academic_artifact_quality_issues

    thesis = tmp_path / "Arforge_Agent_毕业论文.docx"
    proposal = tmp_path / "Arforge_Agent_开题报告.docx"
    thesis_pdf = tmp_path / "Arforge_Agent_毕业论文.pdf"
    proposal_pdf = tmp_path / "Arforge_Agent_开题报告.pdf"
    _write_academic_docx(thesis, proposal=False, valid=True)
    _write_academic_docx(proposal, proposal=True, valid=True)
    _write_page_count_pdf(thesis_pdf, 5)
    _write_page_count_pdf(proposal_pdf, 2)
    calls = [
        {"success": True, "tool": "open_file", "args": {"path": str(thesis)}},
        {"success": True, "tool": "open_file", "args": {"path": str(proposal)}},
        {"success": True, "tool": "open_file", "args": {"path": str(thesis_pdf)}},
        {"success": True, "tool": "open_file", "args": {"path": str(proposal_pdf)}},
    ]

    issues = _academic_artifact_quality_issues(
        "按标准论文格式完成论文和开题报告", calls
    )

    assert any("thesis PDF has only 5 page(s)" in issue for issue in issues)
    assert any("proposal PDF has only 2 page(s)" in issue for issue in issues)


def test_academic_quality_gate_accepts_substantial_derived_pdfs(tmp_path):
    from arforge.core.agentic import _academic_artifact_quality_issues

    thesis = tmp_path / "Arforge_Agent_毕业论文.docx"
    proposal = tmp_path / "Arforge_Agent_开题报告.docx"
    thesis_pdf = tmp_path / "Arforge_Agent_毕业论文.pdf"
    proposal_pdf = tmp_path / "Arforge_Agent_开题报告.pdf"
    _write_academic_docx(thesis, proposal=False, valid=True)
    _write_academic_docx(proposal, proposal=True, valid=True)
    _write_page_count_pdf(thesis_pdf, 10)
    _write_page_count_pdf(proposal_pdf, 3)
    calls = [
        {"success": True, "tool": "open_file", "args": {"path": str(thesis)}},
        {"success": True, "tool": "open_file", "args": {"path": str(proposal)}},
        {"success": True, "tool": "open_file", "args": {"path": str(thesis_pdf)}},
        {"success": True, "tool": "open_file", "args": {"path": str(proposal_pdf)}},
    ]

    assert _academic_artifact_quality_issues(
        "按标准论文格式完成论文和开题报告", calls
    ) == []


def test_academic_quality_gate_rejects_raw_markdown_in_rendered_pdf(tmp_path):
    from arforge.core.agentic import _academic_artifact_quality_issues

    thesis = tmp_path / "Arforge_Agent_毕业论文.docx"
    proposal = tmp_path / "Arforge_Agent_开题报告.docx"
    thesis_pdf = tmp_path / "Arforge_Agent_毕业论文.pdf"
    proposal_pdf = tmp_path / "Arforge_Agent_开题报告.pdf"
    _write_academic_docx(thesis, proposal=False, valid=True)
    _write_academic_docx(proposal, proposal=True, valid=True)
    _write_page_count_pdf(thesis_pdf, 10)
    _write_page_count_pdf(proposal_pdf, 3)
    payload = proposal_pdf.read_bytes()
    raw_rule = ("\ufeff---").encode("utf-16-be").hex().upper().encode("ascii")
    proposal_pdf.write_bytes(
        payload.replace(b"%%EOF", b"<" + raw_rule + b"> Tj\n%%EOF")
    )
    calls = [
        {"success": True, "tool": "open_file", "args": {"path": str(thesis)}},
        {"success": True, "tool": "open_file", "args": {"path": str(proposal)}},
        {"success": True, "tool": "open_file", "args": {"path": str(thesis_pdf)}},
        {"success": True, "tool": "open_file", "args": {"path": str(proposal_pdf)}},
    ]

    issues = _academic_artifact_quality_issues(
        "按标准论文格式完成论文和开题报告", calls
    )

    assert any(
        "raw Markdown/HTML source markup (---)" in issue for issue in issues
    )


def test_academic_pdf_gate_rejects_ascii_blockquote_marker(tmp_path):
    from arforge.core.agentic import _pdf_academic_quality_issues

    proposal_pdf = tmp_path / "开题报告.pdf"
    _write_page_count_pdf(proposal_pdf, 3)
    payload = proposal_pdf.read_bytes()
    proposal_pdf.write_bytes(payload.replace(b"%%EOF", b"(>) Tj\n%%EOF"))

    issues = _pdf_academic_quality_issues(proposal_pdf, proposal=True)

    assert any("raw Markdown/HTML source markup (>)" in issue for issue in issues)


def test_academic_quality_gate_validates_existing_pdf_paired_with_read_docx(tmp_path):
    from arforge.core.agentic import _academic_artifact_quality_issues

    proposal = tmp_path / "Arforge_Agent_\u5f00\u9898\u62a5\u544a.docx"
    proposal_pdf = proposal.with_suffix(".pdf")
    _write_academic_docx(proposal, proposal=True, valid=True)
    _write_page_count_pdf(proposal_pdf, 3)
    payload = proposal_pdf.read_bytes()
    proposal_pdf.write_bytes(payload.replace(b"%%EOF", b"(>) Tj\n%%EOF"))
    calls = [
        {
            "success": True,
            "tool": "read_office_document",
            "args": {"path": str(proposal)},
        }
    ]

    issues = _academic_artifact_quality_issues(
        "Create a PDF \u5f00\u9898\u62a5\u544a in standard academic format",
        calls,
    )

    assert any("raw Markdown/HTML source markup (>)" in issue for issue in issues)


def test_academic_docx_quality_gate_rejects_letter_placeholders(tmp_path):
    from arforge.core.agentic import _academic_artifact_quality_issues

    thesis = tmp_path / "Arforge_Agent_毕业论文.docx"
    proposal = tmp_path / "Arforge_Agent_开题报告.docx"
    _write_academic_docx(thesis, proposal=False, valid=False)
    _write_academic_docx(proposal, proposal=True, valid=False)
    calls = [
        {"success": True, "tool": "read_office_document", "args": {"path": str(thesis)}},
        {"success": True, "tool": "read_office_document", "args": {"path": str(proposal)}},
    ]

    issues = _academic_artifact_quality_issues(
        "按标准论文格式完成论文和开题报告", calls
    )

    assert any("not A4" in issue for issue in issues)
    assert any("manual-update placeholder" in issue for issue in issues)
    assert any("companion Markdown" in issue for issue in issues)
    assert any("populated table of contents" in issue for issue in issues)


def test_academic_docx_quality_gate_rejects_unfinished_fields_and_raw_markup(tmp_path):
    from docx import Document
    from docx.shared import Mm

    from arforge.core.agentic import _docx_academic_quality_issues

    thesis = tmp_path / "完整论文.docx"
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    document.add_paragraph("学生姓名：（请填写）")
    document.add_paragraph("<!--pagebreak-->")
    document.add_heading("中文摘要", level=1)
    document.add_heading("Abstract", level=1)
    document.add_heading("目录", level=1)
    for index in range(1, 7):
        document.add_paragraph(f"第{index}章 章节标题\t{index + 3}")
    document.add_heading("第一章 绪论", level=1)
    document.add_heading("参考文献", level=1)
    for index in range(1, 7):
        document.add_paragraph(f"[{index}] Author. Reference {index}. 2026.")
    document.save(thesis)

    issues = _docx_academic_quality_issues(thesis, proposal=False)

    assert any('exactly "摘要"' in issue for issue in issues)
    assert any("unfinished content placeholder" in issue for issue in issues)
    assert any("raw Markdown/HTML source markup" in issue for issue in issues)


def test_academic_docx_quality_gate_rejects_named_template_placeholders(tmp_path):
    from docx import Document
    from docx.shared import Mm

    from arforge.core.agentic import _docx_academic_quality_issues

    proposal = tmp_path / "Arforge_Agent_开题报告.docx"
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    document.add_heading("开题报告", level=1)
    document.add_paragraph("学生姓名：（填写姓名）")
    document.add_paragraph("指导教师：（填写导师姓名）")
    for title in (
        "一、课题背景与意义",
        "二、研究现状",
        "三、研究目标",
        "四、研究内容",
        "五、技术路线",
        "六、创新点",
        "七、进度计划",
        "八、可行性与风险",
        "九、参考文献",
    ):
        document.add_heading(title, level=1)
        document.add_paragraph("完整内容")
    for index in range(1, 7):
        document.add_paragraph(f"[{index}] Author. Reference {index}. 2026.")
    document.save(proposal)

    issues = _docx_academic_quality_issues(proposal, proposal=True)

    assert any(
        "unfinished content placeholder (（填写姓名）)" in issue
        for issue in issues
    )


def test_academic_identity_gate_rejects_invented_personal_data(tmp_path):
    from docx import Document

    from arforge.core.agentic import _docx_unprovided_identity_issues

    thesis = tmp_path / "毕业论文.docx"
    document = Document()
    document.add_paragraph("学生姓名：张明")
    document.add_paragraph("学号：202406001")
    document.add_paragraph("指导教师：王强 教授")
    document.save(thesis)

    issues = _docx_unprovided_identity_issues(
        thesis,
        "请根据项目源码生成完整毕业论文，但不要编造项目事实。",
    )

    assert len(issues) == 3
    assert all("not provided by the user" in issue for issue in issues)


def test_academic_identity_gate_accepts_user_values_or_anonymous_sample(tmp_path):
    from docx import Document

    from arforge.core.agentic import _docx_unprovided_identity_issues

    thesis = tmp_path / "毕业论文.docx"
    document = Document()
    document.add_paragraph("学生姓名：张明")
    document.add_paragraph("学号：——")
    document.add_paragraph("指导教师：不适用")
    document.save(thesis)

    assert _docx_unprovided_identity_issues(
        thesis,
        "学生姓名使用张明，并生成匿名技术研究样稿。",
    ) == []


def test_academic_gate_rejects_docx_that_omits_most_markdown_source(tmp_path):
    from docx import Document

    from arforge.core.agentic import _docx_markdown_source_coverage_issues

    source = tmp_path / "src"
    source.mkdir()
    (source / "完整论文.md").write_text(
        "# 完整论文\n\n" + ("这是完整论文正文内容。" * 500),
        encoding="utf-8",
    )
    thesis = tmp_path / "毕业论文.docx"
    document = Document()
    document.add_heading("第一章 绪论", level=1)
    document.add_paragraph("这里只生成了很短的一段。")
    document.save(thesis)

    issues = _docx_markdown_source_coverage_issues(
        thesis,
        proposal=False,
        task="交付完整论文 DOCX、PDF 和可追溯的 Markdown 源稿。",
    )

    assert any("Regenerate the DOCX and PDF from the complete source" in issue for issue in issues)


def test_academic_gate_rejects_turn_helpers_left_in_protected_workspace(
    monkeypatch, tmp_path
):
    from arforge.core.agentic import _academic_artifact_quality_issues
    from arforge.tools import file_tools

    monkeypatch.setattr(file_tools, "get_workspace", lambda: tmp_path)
    calls = [
        {
            "tool": "write_file",
            "args": {"filename": "gen_docs.py"},
            "success": True,
        }
    ]

    issues = _academic_artifact_quality_issues(
        "按标准论文格式完成论文和开题报告，但不要修改当前工作区",
        calls,
    )

    assert any("temporary writes remain there: gen_docs.py" in issue for issue in issues)


def test_academic_gate_accepts_protected_workspace_helper_after_cleanup(
    monkeypatch, tmp_path
):
    from arforge.core.agentic import _academic_artifact_quality_issues
    from arforge.tools import file_tools

    monkeypatch.setattr(file_tools, "get_workspace", lambda: tmp_path)
    calls = [
        {
            "tool": "write_file",
            "args": {"filename": "gen_docs.py"},
            "success": True,
        },
        {
            "tool": "delete_file",
            "args": {"filename": "gen_docs.py"},
            "success": True,
        },
    ]

    issues = _academic_artifact_quality_issues(
        "按标准论文格式完成论文和开题报告，但不要修改当前工作区",
        calls,
    )

    assert not any("temporary writes remain there" in issue for issue in issues)
